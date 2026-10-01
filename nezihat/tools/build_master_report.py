from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PROJE_TEKNIK_RAPORU.md"
OUTPUT = ROOT / "CineSum_Audio_Master_Proje_Raporu.docx"

NAVY = "17324D"
BLUE = "2E74B5"
PALE = "E8EEF5"
LIGHT = "F2F4F7"
INK = "202A33"
MUTED = "637083"
WHITE = "FFFFFF"
RED = "9B1C1C"


def font(run, size=None, bold=None, color=INK, name="Calibri", italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_fill(cell, PALE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(str(value)), 9, True, NAVY)
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(value)), 9, False, INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_field(paragraph, instruction, display):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_page_number(paragraph):
    add_field(paragraph, " PAGE ", "1")


def add_callout(doc, label, text, fill=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(label + "  "), 10, True, NAVY)
    font(p.add_run(text), 10, False, INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 10.5)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    font(p.add_run(text), 10.5)


def add_code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F7F8FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    font(p.add_run(text.rstrip()), 8.5, False, INK, "Consolas")
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def clean_inline(text):
    return re.sub(r"`([^`]+)`", r"\1", text.strip())


def import_markdown_body(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code = []
    table_rows = []
    para = []

    def flush_para():
        nonlocal para
        if para:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            font(p.add_run(clean_inline(" ".join(para))), 10.5)
            para = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            rows = [[clean_inline(x) for x in re.split(r"\s*\|\s*", row.strip("| "))] for row in table_rows]
            if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", x) for x in rows[1]):
                rows.pop(1)
            cols = max(len(r) for r in rows)
            widths = [9360 // cols] * cols
            widths[-1] += 9360 - sum(widths)
            add_table(doc, rows[0], [r + [""] * (cols - len(r)) for r in rows[1:]], widths)
            table_rows = []

    started = False
    for line in lines:
        if line.startswith("# ") and not started:
            continue
        if line.startswith("**Rapor tarihi:") or line.startswith("**Proje dizini:") or line.startswith("**Ana çalışma"):
            continue
        if line.startswith("```"):
            flush_para(); flush_table()
            if in_code:
                add_code(doc, "\n".join(code)); code = []; in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code.append(line); continue
        if line.startswith("|"):
            flush_para(); table_rows.append(line); continue
        flush_table()
        m = re.match(r"^(#{2,4})\s+(.*)", line)
        if m:
            flush_para(); started = True
            level = min(len(m.group(1)) - 1, 3)
            doc.add_heading(clean_inline(m.group(2)), level=level)
        elif re.match(r"^\d+\.\s+", line):
            flush_para(); add_number(doc, re.sub(r"^\d+\.\s+", "", clean_inline(line)))
        elif line.startswith("- "):
            flush_para(); add_bullet(doc, clean_inline(line[2:]))
        elif not line.strip():
            flush_para()
        else:
            para.append(line.strip())
    flush_para(); flush_table()


def setup_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(hp.add_run("CineSum Audio | Master Proje Raporu"), 8.5, False, MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(fp.add_run("CineSum Audio  •  Ekip içi teknik referans  •  "), 8, False, MUTED)
    add_page_number(fp)


def main():
    doc = Document()
    setup_styles(doc)
    add_header_footer(doc.sections[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("CINESUM AUDIO"), 12, True, BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("MASTER PROJE RAPORU"), 28, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    font(p.add_run("Mimari, çalışma akışı, algoritmalar, çıktı sözleşmeleri, testler ve ekip devir kılavuzu"), 14, False, MUTED)
    add_callout(doc, "Belgenin amacı", "Projeye yeni katılan bir ekip üyesinin kodu çalıştırabilmesi, çıktıları yorumlayabilmesi, teknik kararların gerekçesini anlayabilmesi ve güvenli biçimde geliştirmeye devam edebilmesi için tek kaynak oluşturmaktır.", PALE)
    add_table(doc, ["Belge", "Değer"], [
        ["Sürüm", "1.0 - Master rapor"],
        ["Tarih", "15 Ağustos 2026"],
        ["Kapsam", "cinesum-audio çalışma alanının mevcut durumu"],
        ["Birincil giriş noktası", "process_video.py"],
        ["Hedef okuyucu", "Yazılım, veri/ML, test ve ürün ekipleri"],
        ["Doğrulama", "6/6 otomatik test başarılı (.venv, 15 Ağustos 2026)"],
    ], [2400, 6960])
    doc.add_page_break()

    doc.add_heading("Belgeyi nasıl kullanmalı?", 1)
    add_bullet(doc, "Hızlı yöneticilik özeti için Bölüm 1-3'ü okuyun.")
    add_bullet(doc, "Sistemi çalıştıracak ekip üyeleri Bölüm 4-6 ve operasyon kılavuzuna odaklansın.")
    add_bullet(doc, "Algoritma veya model davranışını değiştirecek geliştiriciler mevcut teknik raporun ayrıntılı akışını ve karar kayıtlarını birlikte okusun.")
    add_bullet(doc, "Hata ayıklama, teslim ve devir için kontrol listeleri ile çıktı sözleşmelerini kullanın.")
    doc.add_heading("İçindekiler", 1)
    p = doc.add_paragraph()
    add_field(p, ' TOC \\o "1-3" \\h \\z \\u ', "İçindekiler alanını Word'de sağ tıklayıp 'Alanı Güncelleştir' seçeneğiyle yenileyin.")
    doc.add_page_break()

    doc.add_heading("1. Yönetici özeti", 1)
    doc.add_paragraph("CineSum Audio, video sesini tek bir komutla analiz eden Python tabanlı bir işlem hattıdır. Sistem konuşmayı metne dönüştürür, konuşmacı turlarını ayırır, kelimeleri konuşmacılara atar, eş zamanlı konuşmaları ve isteğe bağlı dış ses olaylarını raporlar. Çıktılar hem insan tarafından okunabilir metin hem de makine tarafından işlenebilir JSON olarak saklanır.")
    add_callout(doc, "Mevcut durum", "Ana akış çalışır durumdadır; altı davranış testi geçmektedir. Depoda iki örnek çalışmanın sonuçları vardır. En önemli üretim riskleri CPU performansı, model indirme/erişim bağımlılığı, otomatik diarization hataları ve transkripsiyon/diarization önbelleğinin henüz bulunmamasıdır.")
    add_table(doc, ["Alan", "Durum", "Ekip için anlamı"], [
        ["Ana işlem hattı", "Çalışır", "process_video.py tek giriş noktasıdır."],
        ["Konuşmacı düzeltmesi", "Çalışır", "Embedding tabanlı otomatik düzeltme ve video bazlı manuel override birlikte kullanılabilir."],
        ["Testler", "6/6 başarılı", "Sınır ataması, override, merge ve overlap davranışları korunmaktadır."],
        ["Örnek çıktı", "Mevcut", "output/test ve output/test3 klasörleri inceleme için kullanılabilir."],
        ["GPU", "Mevcut ortam CPU", "Uzun videolarda işlem süresi yüksek olabilir."],
        ["Önbellek", "Kısmi", "Yalnızca dış ses olaylarında vardır; ASR ve diarization yeniden çalışır."],
    ], [2100, 1500, 5760])

    doc.add_heading("2. Sistem bağlamı ve sınırlar", 1)
    doc.add_paragraph("Sistem bir video dosyasını girdi olarak alır ve yerel dosya sistemine video bazlı bir sonuç paketi yazar. FFmpeg medya çözümleme katmanıdır; Hugging Face üzerinden edinilen modeller transkripsiyon, diarization, embedding ve dış ses sınıflandırmasını gerçekleştirir. Sistem gerçek kişi kimliği üretmez; SPEAKER_XX etiketleri yalnızca aynı video içindeki kümelerdir.")
    add_code(doc, "Video → FFmpeg/PCM → Pyannote diarization → Embedding düzeltmesi\n      ↘ Faster-Whisper ASR → Kelime/konuşmacı eşleme → Birleştirme\n      ↘ AST dış ses analizi (isteğe bağlı)\n      → result.json + result.txt + ara/denetim çıktıları")
    doc.add_heading("2.1. Kapsam dışı noktalar", 2)
    for item in [
        "Gerçek kişi adının otomatik belirlenmesi veya yüz tanıma.",
        "Çıktıların web arayüzünde görsel düzenlenmesi.",
        "Dağıtık iş kuyruğu, çok kullanıcılı servis veya bulut dağıtımı.",
        "Diarization ve transkripsiyon için güvenli yeniden kullanım önbelleği.",
        "Model doğruluğunun akademik veri kümesi üzerinde nicel kıyaslaması.",
    ]: add_bullet(doc, item)

    doc.add_heading("3. Ekip sorumluluk haritası", 1)
    add_table(doc, ["Rol", "Ana sorumluluk", "Kontrol noktaları"], [
        ["Teknik lider", "Mimari kararlar ve sürüm kapsamı", "Model/algoritma değişiklikleri, performans-risk dengesi"],
        ["ML geliştiricisi", "Diarization, embedding, ASR ve eşikler", "Ham/düzeltilmiş sonuç karşılaştırması, regresyon örnekleri"],
        ["Uygulama geliştiricisi", "CLI, dosya akışı, hata yönetimi", "Girdi doğrulama, deterministik çıktı şeması"],
        ["Test sorumlusu", "Davranış ve uçtan uca doğrulama", "Birim testleri, örnek video kabul kriterleri"],
        ["Ürün/içerik sorumlusu", "Kullanım senaryosu ve çıktı kalitesi", "Metin okunabilirliği, gerçek konuşmacı geri bildirimi"],
        ["Operasyon", "Kurulum, model erişimi, FFmpeg/GPU", "Sürüm sabitleme, disk/CPU/GPU ve token erişimi"],
    ], [1600, 3500, 4260])

    doc.add_heading("4. Hızlı başlangıç ve işletim kılavuzu", 1)
    for step in [
        "Python 3.13 sanal ortamını etkinleştirin ve requirements.txt bağımlılıklarını kurun.",
        "FFmpeg'in sistem PATH'inde, FFMPEG_PATH değişkeninde veya proje içindeki beklenen konumda erişilebilir olduğunu doğrulayın.",
        "Hugging Face hesabında pyannote model koşullarını kabul edin ve hf auth login ile oturum açın.",
        "Videoyu input klasörüne koyun veya mutlak/göreli yolunu komutta verin.",
        "İlk doğrulama için dış ses analizini kapatarak çalıştırın: python process_video.py input\\video.mp4 --skip-events.",
        "output/<video_adı>/result.txt ve result.json dosyalarını; ardından diarization_raw.json ile diarization.json farkını kontrol edin.",
    ]: add_number(doc, step)
    add_callout(doc, "Önerilen varsayılan", "Doğruluk öncelikliyse medium Whisper + beam size 5 korunmalı, yalnızca gerekli değilse --skip-events kullanılmalıdır. --fast, hız için doğruluk ödünüdür.")

    doc.add_heading("5. Komut satırı sözleşmesi", 1)
    cli_rows = [
        ["video", "zorunlu", "-", "İşlenecek video yolu"],
        ["--language", "isteğe bağlı", "en", "Whisper konuşma dili"],
        ["--whisper-model", "isteğe bağlı", "medium", "ASR modeli; --fast ile varsayılan small olur"],
        ["--device", "isteğe bağlı", "auto", "auto/cpu/cuda"],
        ["--fast", "bayrak", "kapalı", "Daha küçük ASR ve daha seyrek olay taraması"],
        ["--skip-events", "bayrak", "kapalı", "AST dış ses analizini atlar"],
        ["--recompute-events", "bayrak", "kapalı", "Olay önbelleğini yok sayar"],
        ["--event-threshold", "isteğe bağlı", "0.35", "Dış ses kabul eşiği (0-1)"],
        ["--speakers", "isteğe bağlı", "otomatik", "Kesin konuşmacı sayısı"],
        ["--min-speakers / --max-speakers", "isteğe bağlı", "otomatik", "Küme sayısı aralığı"],
        ["--speaker-threshold", "isteğe bağlı", "0.40", "Pyannote clustering hassasiyeti"],
        ["--no-speaker-correction", "bayrak", "kapalı", "Embedding düzeltmesini kapatır"],
        ["--disable-intra-speaker-split", "bayrak", "kapalı", "Aynı etiket içindeki ayrıştırmayı kapatır"],
        ["--speaker-merge-threshold", "isteğe bağlı", "0.88", "Benzer kümeleri birleştirme eşiği"],
        ["--reassign-similarity", "isteğe bağlı", "0.72", "Yeniden atama asgari benzerliği"],
        ["--reassign-margin", "isteğe bağlı", "0.15", "Yeniden atama fark şartı"],
        ["--speaker-overrides", "isteğe bağlı", "video klasörü", "Manuel düzeltme JSON yolu"],
    ]
    add_table(doc, ["Seçenek", "Tür", "Varsayılan", "Etkisi"], cli_rows, [2350, 1400, 1450, 4160])

    doc.add_heading("6. Çıktı ve veri sözleşmeleri", 1)
    add_table(doc, ["Dosya", "Kaynak", "Tüketim amacı"], [
        ["result.txt", "Nihai timeline", "İnsan incelemesi ve hızlı kalite kontrol"],
        ["result.json", "Tüm aşamalar", "Uygulama entegrasyonu ve arşiv"],
        ["transcript.json", "Faster-Whisper", "Segment/kelime zamanları ve ASR denetimi"],
        ["diarization_raw.json", "Pyannote", "Düzeltme öncesi referans"],
        ["diarization.json", "Düzeltme katmanı", "Nihai turlar, overlap ve rapor özeti"],
        ["speaker_correction_report.json", "Embedding katmanı", "Birleştirme, bölme ve yeniden atama denetimi"],
        ["speaker_overrides.json", "Kullanıcı", "Video bazlı kesin düzeltmeler"],
        ["sound_events.json", "AST", "Olaylar ve önbellek anahtarı"],
    ], [2500, 2200, 4660])
    doc.add_heading("6.1. result.json alanları", 2)
    add_table(doc, ["Alan", "Anlam"], [
        ["input_file", "Kaynak video yolu"], ["language", "Transkripsiyon dili"],
        ["speaker_count", "Düzeltme sonrası benzersiz konuşmacı sayısı"],
        ["segments", "Konuşma blokları: start, end, speaker, text, type"],
        ["sound_events", "Dış ses olayları; skip-events durumunda boş olabilir"],
        ["overlaps", "Aynı zaman aralığındaki çoklu konuşmacılar"],
        ["speaker_correction", "Diarization düzeltme özeti"],
        ["timeline", "Konuşma, overlap ve olayların zamana göre birleşik görünümü"],
    ], [2500, 6860])

    doc.add_heading("7. Doğrulama ve kabul ölçütleri", 1)
    doc.add_paragraph("15 Ağustos 2026 tarihinde proje içindeki .venv kullanılarak altı test yeniden çalıştırılmış ve tamamı geçmiştir. Bu testler hızlıdır ve model indirmeden saf eşleme/düzeltme davranışlarını doğrular; uçtan uca model doğruluğunu ölçmez.")
    add_table(doc, ["Test davranışı", "Sonuç"], [
        ["Kelime zamanı yoksa segment metnini koruma", "Başarılı"],
        ["En büyük örtüşmeye göre kelime atama", "Başarılı"],
        ["Manuel override ile segment bölme", "Başarılı"],
        ["Aynı konuşmacının ardışık Whisper segmentlerini birleştirme", "Başarılı"],
        ["Overlap için düzeltilmiş en yakın kimliği kullanma", "Başarılı"],
        ["Cümle başı zaman kaymasında tek kelimelik yanlış etiketi önleme", "Başarılı"],
    ], [7600, 1760])
    doc.add_heading("7.1. Uçtan uca kabul kontrol listesi", 2)
    for item in [
        "Komut hatasız tamamlandı ve beklenen video klasörü oluştu.",
        "result.txt kronolojik, okunabilir ve boş değil.",
        "speaker_count ile segmentlerdeki benzersiz konuşmacı sayısı uyumlu.",
        "Diarization düzeltmeleri raporda açıklanıyor; ham veri korunuyor.",
        "Overlap etiketleri düzeltilmiş konuşmacı kimlikleriyle uyumlu.",
        "Manuel override kullanıldıysa kapsam yalnızca ilgili videoya ait.",
        "Eşik/model seçenekleri teslim notunda kayıtlı.",
    ]: add_bullet(doc, "☐ " + item)

    doc.add_heading("8. Örnek çalışmaların anlık görünümü", 1)
    add_table(doc, ["Örnek", "Konuşmacı", "Konuşma bloğu", "Overlap", "Timeline"], [
        ["output/test", "5", "12", "2", "14"],
        ["output/test3", "4", "36", "12", "48"],
    ], [2800, 1500, 2000, 1400, 1660])
    add_callout(doc, "Yorumlama notu", "Bu sayılar doğruluk metriği değildir; yalnızca depodaki sonuç paketlerinin yapısal özetidir. Model kalitesi, videoyu dinleyerek/izleyerek referans anotasyonla karşılaştırılmalıdır.")

    doc.add_heading("9. Riskler, sınırlamalar ve azaltma planı", 1)
    add_table(doc, ["Risk", "Etkisi", "Azaltma"], [
        ["CPU üzerinde uzun çalışma", "Teslim süresi ve kullanıcı deneyimi", "GPU kurulumu, olay analizini kapatma, güvenli ASR/diarization önbelleği"],
        ["Model ağı/token bağımlılığı", "İlk kurulumda başarısızlık", "Kurulum ön kontrolü, erişim koşullarının belgelenmesi, yerel model cache doğrulaması"],
        ["Kısa/benzer seslerde diarization hatası", "Yanlış konuşmacı", "Embedding doğrulaması, override, bilinen kişi sayısı, örnek bazlı QA"],
        ["Overlap ve müzik/gürültü", "Sınır ve kimlik belirsizliği", "Ham/nihai sonuç ayrımı, güven puanı ve manuel inceleme"],
        ["Şema değişikliği", "Tüketici uygulamaların kırılması", "Şema sürümü, JSON sözleşme testi ve geriye uyumluluk"],
        ["Testlerin model kalitesini ölçmemesi", "Yanlış güven", "Referans anotasyonlu uçtan uca değerlendirme seti"],
    ], [2600, 2750, 4010])

    doc.add_heading("10. Önceliklendirilmiş yol haritası", 1)
    add_table(doc, ["Öncelik", "İş", "Tamamlanma ölçütü"], [
        ["P0", "Transkripsiyon ve diarization önbelleği", "Video özeti + ayar/model sürümüyle güvenli cache invalidation"],
        ["P0", "Uçtan uca örnek video regresyonu", "Referans anotasyon ve ölçülen WER/DER benzeri metrikler"],
        ["P1", "JSON şema sürümleme", "schema_version ve sözleşme testleri"],
        ["P1", "Konuşmacı isim eşleme", "SPEAKER_XX → kullanıcı onaylı görünen ad"],
        ["P1", "Güven puanları", "Düşük güvenli ASR/atama/olayların işaretlenmesi"],
        ["P2", "Görsel düzeltme arayüzü", "Timeline üzerinde dinle-düzelt-kaydet akışı"],
        ["P2", "Dağıtım/paketleme", "Tekrarlanabilir kurulum ve ortam sağlık kontrolü"],
    ], [1200, 3300, 4860])

    doc.add_heading("11. Devir teslim kontrol listesi", 1)
    for item in [
        "Yeni ekip üyesi README ve bu master raporu okudu.",
        "process_video.py --help çıktısı ve en az bir örnek komut görüldü.",
        "Altı birim test yerel .venv ile çalıştırıldı.",
        "output/test paketinde raw/final diarization ve result dosyaları karşılaştırıldı.",
        "speaker_overrides.json kapsamının video bazlı olduğu anlaşıldı.",
        "Model/token/FFmpeg bağımlılıkları ekip ortamında doğrulandı.",
        "Planlanan değişiklik için etkilenen fonksiyonlar ve çıktı şemaları belirlendi.",
        "Yeni davranış için test ve örnek çıktı güncelleme sorumlusu atandı.",
    ]: add_bullet(doc, "☐ " + item)

    doc.add_page_break()
    doc.add_heading("12. Ayrıntılı teknik çalışma raporu", 1)
    doc.add_paragraph("Aşağıdaki bölüm, 14 Ağustos 2026 tarihli teknik çalışma raporunun içeriklerini master belgenin parçası olarak korur. Başlık sıralaması ve teknik ayrıntılar ekip içi izlenebilirlik amacıyla muhafaza edilmiştir.")
    import_markdown_body(doc)

    doc.add_heading("Ek A. Kaynak kod sorumluluk matrisi", 1)
    add_table(doc, ["Dosya", "Rol", "Değişiklikte dikkat"], [
        ["process_video.py", "Ana uçtan uca orkestrasyon", "CLI geriye uyumluluğu, çıktı şeması, model yükleme maliyeti"],
        ["speaker_correction.py", "Embedding tabanlı yeniden atama/birleştirme/bölme", "Eşiklerin veriyle kalibrasyonu, kısa segment güvenilirliği"],
        ["tests/test_process_video.py", "Saf davranış regresyonları", "Her hata düzeltmesine karşı test ekleme"],
        ["transcribe.py", "Eski deneysel ASR akışı", "Ana üretim yolu değildir"],
        ["diarize.py", "Eski deneysel diarization akışı", "Ana üretim yolu değildir"],
        ["merge_results.py", "Eski deneysel birleştirme", "Ana algoritmayla karıştırılmamalı"],
        ["requirements.txt", "Bağımlılık sürümleri", "Model/kütüphane uyumu ve Python sürümü"],
    ], [2500, 3300, 3560])

    doc.add_heading("Ek B. Kritik fonksiyonlar", 1)
    add_table(doc, ["Fonksiyon", "Sorumluluk"], [
        ["find_ffmpeg", "FFmpeg yürütülebilir dosyasını öncelik sırasıyla bulur"],
        ["load_audio", "Videoyu 16 kHz mono PCM olarak belleğe alır"],
        ["diarize", "Pyannote konuşmacı turlarını üretir"],
        ["correct_diarization", "Embedding tabanlı düzeltme zincirini yürütür"],
        ["transcribe", "Faster-Whisper segment ve kelimelerini üretir"],
        ["assign_word_speakers", "Dinamik programlama ile kelime-konuşmacı dizisini seçer"],
        ["build_speaker_transcript", "Atamaları ardışık konuşma bloklarına dönüştürür"],
        ["update_overlap_speakers", "Overlap kimliklerini düzeltilmiş turlarla eşler"],
        ["detect_events", "AST ile seçili dış ses olaylarını tarar"],
        ["save_outputs", "Nihai JSON ve metin paketini yazar"],
    ], [3300, 6060])

    doc.add_heading("Ek C. Karar kayıtları", 1)
    add_table(doc, ["Karar", "Gerekçe", "Sonuç"], [
        ["Tek giriş noktası process_video.py", "Üç eski betikli akışın yarattığı belirsizliği kaldırmak", "Kurulum ve kullanım sadeleşti"],
        ["Sesin FFmpeg ile belleğe alınması", "Codec/TorchCodec bağımlılığını azaltmak", "Tek PCM kaynağı ve daha öngörülebilir giriş"],
        ["Ham diarization çıktısını korumak", "Otomatik düzeltmelerin denetlenebilirliği", "Önce/sonra karşılaştırması mümkün"],
        ["Video bazlı override", "Konuşmacı etiketlerinin videolar arasında sabit olmaması", "Düzeltme kapsamı güvenli biçimde yerel"],
        ["Viterbi benzeri kelime ataması", "Tek kelimelik sınır sıçramalarını azaltmak", "Cümle içi kimlik sürekliliği"],
        ["Dış ses analizinin isteğe bağlı olması", "CPU maliyetini kontrol etmek", "--skip-events ile güvenli hız kazanımı"],
    ], [2800, 3600, 2960])

    doc.add_heading("Ek D. Sorun giderme", 1)
    add_table(doc, ["Belirti", "Olası neden", "İlk kontrol"], [
        ["FFmpeg bulunamadı", "PATH/FFMPEG_PATH yanlış", "ffmpeg.exe konumunu ve ortam değişkenini doğrulayın"],
        ["Model indirilemiyor", "Ağ, proxy, token veya model koşulları", "hf auth login ve model erişim onayını kontrol edin"],
        ["CUDA seçilemiyor", "CPU PyTorch veya uyumsuz sürücü", "torch.cuda.is_available ve PyTorch/CUDA eşleşmesini doğrulayın"],
        ["Yanlış konuşmacı", "Kısa segment, benzer ses, yanlış kişi sayısı", "raw/final diarization, --speakers ve override kullanımı"],
        ["İlk kelime yanlış kişide", "Zaman damgası sınır kayması", "Varsayılan modu kullanın; ilgili regresyon testini çalıştırın"],
        ["İşlem çok yavaş", "CPU, medium model, olay analizi", "--skip-events; gerekirse --fast ve doğruluk kontrolü"],
        ["Olay sonucu eski", "Önbellek kullanıldı", "--recompute-events ile yeniden tarayın"],
    ], [2600, 3450, 3310])

    doc.add_heading("Ek E. Sürümleme ve güncelleme protokolü", 1)
    for step in [
        "Davranış değişikliğini ve etkilediği çıktı alanlarını tanımlayın.",
        "Mevcut örnek üzerinde önce/sonra sonuçlarını saklayın.",
        "En az bir regresyon testi ekleyin veya mevcut testi güncelleyin.",
        "CLI varsayılanı değişiyorsa README, master rapor ve teslim notunu birlikte güncelleyin.",
        "Model veya eşik değişiyorsa kullanılan sürüm/ayarları result.json metadata alanına taşıyın.",
        "Tüm testleri ve bir uçtan uca örneği çalıştırıp doğrulama tarihini kaydedin.",
    ]: add_number(doc, step)

    props = doc.core_properties
    props.title = "CineSum Audio Master Proje Raporu"
    props.subject = "Teknik mimari, işletim, doğrulama ve ekip devir kılavuzu"
    props.author = "CineSum Audio Ekibi"
    props.keywords = "CineSum, audio, diarization, whisper, teknik rapor, handover"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
